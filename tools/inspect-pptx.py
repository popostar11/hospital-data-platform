# -*- coding: utf-8 -*-
"""
检查一个 pptx 的结构与风格，用于「参照现有汇报 PPT 的格式」。

用法：
    python tools/inspect-pptx.py <文件路径> [--full]

输出：
    - 画布尺寸、页数、每页使用的版式
    - 每页的形状清单（类型/位置/尺寸/文本摘要）
    - 表格与图片明细
    - 全篇用到的字体、字号、颜色统计（用于提炼设计规范）
"""
import sys
from collections import Counter

from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0


def inches(v):
    return None if v is None else round(v / EMU, 2)


def walk(shapes, depth=0):
    for sh in shapes:
        yield sh, depth
        if sh.shape_type == 6:  # GROUP
            yield from walk(sh.shapes, depth + 1)


def run_style(sh):
    """取形状里第一个非空 run 的字体信息"""
    if not sh.has_text_frame:
        return None
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip():
                f = r.font
                color = None
                try:
                    if f.color and f.color.type is not None and f.color.rgb is not None:
                        color = str(f.color.rgb)
                except Exception:
                    pass
                return {
                    "font": f.name,
                    "size": f.size.pt if f.size else None,
                    "bold": f.bold,
                    "color": color,
                }
    return None


def main():
    path = sys.argv[1]
    full = "--full" in sys.argv
    prs = Presentation(path)

    print(f"文件：{path}")
    print(f"画布：{inches(prs.slide_width)} x {inches(prs.slide_height)} 英寸"
          f"（比例 {prs.slide_width / prs.slide_height:.3f}）")
    print(f"页数：{len(prs.slides._sldIdLst)}")
    print(f"母版数：{len(prs.slide_masters)}  版式数：{len(prs.slide_layouts)}")
    print()

    fonts, sizes, colors = Counter(), Counter(), Counter()

    for i, slide in enumerate(prs.slides, 1):
        try:
            layout = slide.slide_layout.name
        except Exception:
            layout = "?"
        shapes = list(walk(slide.shapes))
        kinds = Counter(str(sh.shape_type).split(" ")[0] for sh, _ in shapes)
        print(f"── 第 {i} 页  [版式: {layout}]  形状 {len(shapes)} 个  {dict(kinds)}")

        for sh, depth in shapes:
            pad = "   " + "  " * depth
            pos = f"({inches(sh.left)},{inches(sh.top)}) {inches(sh.width)}x{inches(sh.height)}"
            st = run_style(sh)
            if st:
                if st["font"]:
                    fonts[st["font"]] += 1
                if st["size"]:
                    sizes[round(st["size"], 1)] += 1
                if st["color"]:
                    colors[st["color"]] += 1

            if sh.has_text_frame and sh.text_frame.text.strip():
                txt = sh.text_frame.text.replace("\n", " ⏎ ")
                if not full and len(txt) > 72:
                    txt = txt[:72] + "…"
                style = ""
                if st:
                    style = f"  <{st['font']} {st['size']}pt{'' if not st['bold'] else ' B'}{'' if not st['color'] else ' #' + st['color']}>"
                print(f"{pad}· TEXT {pos}{style}")
                print(f"{pad}    {txt}")
            elif sh.shape_type is not None and "TABLE" in str(sh.shape_type):
                tbl = sh.table
                print(f"{pad}· TABLE {pos}  {len(tbl.rows)}行 x {len(tbl.columns)}列")
                for r_i, row in enumerate(tbl.rows):
                    cells = [c.text.replace("\n", " ")[:26] for c in row.cells]
                    print(f"{pad}    [{r_i}] {' | '.join(cells)}")
            elif sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
                print(f"{pad}· PICTURE {pos}")
            elif sh.has_text_frame:
                print(f"{pad}· SHAPE {sh.shape_type} {pos} (无文字)")
            else:
                print(f"{pad}· {sh.shape_type} {pos}")

            if full and sh.has_text_frame and len(sh.text_frame.text) > 72:
                for line in sh.text_frame.text.split("\n"):
                    print(f"{pad}      {line}")
        print()

    print("=" * 70)
    print("字体使用：", dict(fonts.most_common()))
    print("字号使用：", dict(sizes.most_common(14)))
    print("文字颜色：", dict(colors.most_common(14)))


if __name__ == "__main__":
    main()
